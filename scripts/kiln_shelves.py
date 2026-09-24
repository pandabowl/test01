#!/usr/bin/env python3
"""
Kiln shelf catalog cleanup: consistent product titles plus the data behind an
Experro "filter by size" for Sheffield Pottery's kiln shelves.

Subcommands:
    fetch    Pull every kiln shelf (product type "Kiln Furniture") from the
             Shopify Admin API into kiln-shelves/snapshot.json.
    plan     Parse brand, shape, size and thickness out of each shelf in the
             snapshot and write the review sheet kiln-shelves/plan.csv (and
             plan.xlsx): current vs. proposed title and SEO title, the three
             filter values, and missing category tags. Writes nothing to
             Shopify.
    apply    Push the approved rows of plan.csv to Shopify. Dry run unless
             --confirm is passed.
    restore  Put titles and SEO titles back to what snapshot.json recorded.
             Dry run unless --confirm is passed.

Usage:
    export SHOPIFY_STORE_DOMAIN=your-store.myshopify.com
    export SHOPIFY_ADMIN_ACCESS_TOKEN=shpat_xxx   # read_products, write_products
    python3 scripts/kiln_shelves.py fetch
    python3 scripts/kiln_shelves.py plan
    # review / edit kiln-shelves/plan.csv, set approve=no on rows to skip
    python3 scripts/kiln_shelves.py apply              # dry run
    python3 scripts/kiln_shelves.py apply --confirm    # write to Shopify

Filter data model:
    Three product metafields, each a single-line text with a fixed list of
    choices (the same pattern the store already uses for custom.kiln_type,
    custom.voltage, etc.), so Experro can offer them as facets:

        custom.kiln_shelf_size       12" x 24"  |  21" Diameter
        custom.kiln_shelf_shape      Rectangle  |  Half Round  |  Full 10-Sided
        custom.kiln_shelf_thickness  5/16"      |  3/4"

    Rectangles and squares are always written short side first, so Advancer's
    "12 x 24" and Spectrum's "24 x 12" land on the same filter value. Round
    and multi-sided shelves (full or half) are filtered by diameter.

Title pattern:
    <Brand/line> Kiln Shelf - <size> <shape>, <thickness> Thick
    e.g. Advancer Silicon Carbide Kiln Shelf - 12" x 24" Rectangle, 5/16" Thick
         Spectrum High Alumina Kiln Shelf - 21" Half Round, 3/4" Thick

Handles (URLs) are never changed.
"""
import argparse
import csv
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from fractions import Fraction

API_VERSION = os.environ.get("SHOPIFY_API_VERSION", "2026-07")

DEFAULT_SNAPSHOT = "kiln-shelves/snapshot.json"
DEFAULT_PLAN_CSV = "kiln-shelves/plan.csv"
DEFAULT_PLAN_XLSX = "kiln-shelves/plan.xlsx"
DEFAULT_SYNONYMS = "kiln-shelves/experro-synonyms.csv"

STOREFRONT = "https://www.sheffield-pottery.com/products/"
SEO_SUFFIX = " | Sheffield Pottery"
SEO_MAX_LEN = 70
IN = '"'

# Every visible shelf should carry these, so it shows up in the
# "Kiln Shelves and Posts" (tag "Kiln Shelves") and "Kiln Room" smart
# collections that the storefront categories are built from.
CATEGORY_TAGS = ["Kiln Shelves", "Kiln Shelves - Posts - Cones and Kiln Firing Accessories"]
HIDDEN_TAGS = {"options_hidden_product", "hideonstorefront"}

# key -> (definition name, description)
METAFIELDS = {
    "kiln_shelf_shape": (
        "Kiln Shelf Shape",
        "Shelf outline for the kiln shelf size filter: Full/Half Round, Full/Half 8-, 10- or 12-Sided, Rectangle, Square.",
    ),
    "kiln_shelf_size": (
        "Kiln Shelf Size",
        'Kiln shelf size filter. Rectangles/squares short side first (12" x 24"); round and multi-sided shelves by diameter (21" Diameter).',
    ),
    "kiln_shelf_thickness": (
        "Kiln Shelf Thickness",
        'Kiln shelf thickness in inches, e.g. 5/16", 3/4", 1".',
    ),
}
PLAN_FIELD_FOR_KEY = {
    "kiln_shelf_shape": "shelf_shape",
    "kiln_shelf_size": "shelf_size",
    "kiln_shelf_thickness": "shelf_thickness",
}

SHAPE_ORDER = [
    "Full Round", "Half Round",
    "Full 8-Sided", "Half 8-Sided",
    "Full 10-Sided", "Half 10-Sided",
    "Full 12-Sided", "Half 12-Sided",
    "Rectangle", "Square",
]

# Brand / product line: title prefix, SEO title prefix, line collection tag.
LINES = {
    "advancer": ("Advancer Silicon Carbide Kiln Shelf", "Advancer Kiln Shelf", "ADVANCER Kiln Shelves"),
    "crystolon": (
        "Crystolon Oxide-Bonded Silicon Carbide Kiln Shelf",
        "Crystolon Silicon Carbide Kiln Shelf",
        "Silicon Carbide Kiln Shelves",
    ),
    "spectrum": ("Spectrum High Alumina Kiln Shelf", "Spectrum High Alumina Kiln Shelf", "High Alumina Cone 11 Kiln Shelves"),
    "spectrum_semi_hollow": (
        "Spectrum High Alumina Semi-Hollow Kiln Shelf",
        "Spectrum Semi-Hollow Kiln Shelf",
        "High Alumina Cone 11 Kiln Shelves",
    ),
    "corelite": (
        "Cedar Heights CoreLite Semi-Hollow Kiln Shelf",
        "CoreLite Semi-Hollow Kiln Shelf",
        "Semi Hollow CoreLite Kiln Shelves",
    ),
    "gillespie": ("Gillespie Hollow Core Kiln Shelf", "Gillespie Hollow Core Kiln Shelf", None),
    "olympic": ("Olympic Kiln Shelf", "Olympic Kiln Shelf", "Olympic Kiln Shelves and Furniture Kits"),
}

# Per-product corrections and review notes, keyed by product handle. Anything
# set here wins over what the title parser found.
OVERRIDES = {
    "olympic-21-x-1-2-shelf-olhs21": {
        "shape": None,
        "thickness": None,
        "diameter": 21,
        "approve": "no",
        "notes": [
            'Title/description only say "21" x 1/2 shelf": likely a 21" HALF shelf (SKU OLHS21, $59), '
            "thickness and round vs 10-sided unknown. Fill in shelf_shape, shelf_thickness and new_title, then set approve=yes.",
        ],
    },
    "advancer-kiln-shelf-20-full-10-sided-ssfadv20f": {
        "thickness": Fraction(5, 16),
        "notes": ['Thickness not in title; set to Advancer standard 5/16" - confirm.'],
    },
    "advancer-kiln-shelf-20-half-10-sided-ssfadv20h": {
        "thickness": Fraction(5, 16),
        "notes": ['Thickness not in title; set to Advancer standard 5/16" - confirm.'],
    },
    "advancer-kiln-shelf-21-full-10-sided-ssfadv21f": {
        "thickness": Fraction(5, 16),
        "notes": [
            'Thickness not in title; set to Advancer standard 5/16" - confirm.',
            "Vendor part no. (AKS-ADV2020F10S-1) duplicates the 20\" full shelf's.",
        ],
    },
    "advancer-kiln-shelf-21-half-10-sided-ssfadv21h": {
        "thickness": Fraction(5, 16),
        "notes": ['Thickness not in title; set to Advancer standard 5/16" - confirm.'],
    },
    "advancer-kiln-shelf-26-half-12-sided-ssfadv26h": {
        "thickness": Fraction(5, 16),
        "notes": ['Thickness not in title; set to Advancer standard 5/16" - confirm.'],
    },
    "corelite-kiln-semi-hollow-shelves-ch161558": {
        "notes": [
            'Octagon listed as 16" x 15" (corner-to-corner x flat-to-flat?); filed under 16" Diameter - confirm.',
        ],
    },
    "kiln-shelf-spha09": {
        "notes": ['Description says 26" x 13" x 1"; title and SKU say 3/4" (the 1" version is archived). Fix the description.'],
    },
    "advancer-kiln-shelf-13-x26-x-5-16-silicon-carbide-ssfadv1326": {
        "notes": ['Description opens with "12 x28"; fix the description.'],
    },
    "advancer-kiln-shelf-14-x28-x-5-16-silicon-carbide-ssfadv1428": {
        "notes": ['Description opens with "11 x 22"; fix the description.'],
    },
    "16-x-13-x-5-16-advancer®-kiln-shelf": {
        "notes": ["Vendor part no. field (AKS-ADV1326X5/16-I) belongs to the 13 x 26 shelf."],
    },
    "26-1-2-x-1-half-round-high-alumina-kiln-shelf": {
        "notes": ['Old SEO title said 27 1/2" and meta description says 28 1/2"; fix the meta description too.'],
    },
    "kiln-shelf-spha01": {
        "notes": ['Old SEO title said 27 1/2" and meta description says 28 1/2"; fix the meta description too.'],
    },
    "kiln-shelf-spha14": {
        "notes": ['Vendor part no. #20H is shared with the 20" Half 10-Sided shelf.'],
    },
}

PLAN_COLUMNS = [
    "approve",
    "current_title",
    "new_title",
    "shelf_size",
    "shelf_shape",
    "shelf_thickness",
    "current_seo_title",
    "new_seo_title",
    "tags_to_add",
    "notes",
    "status",
    "vendor",
    "sku",
    "handle",
    "product_id",
    "url",
]
COLUMN_WIDTHS = {
    "approve": 9, "current_title": 50, "new_title": 62, "shelf_size": 15, "shelf_shape": 14,
    "shelf_thickness": 10, "current_seo_title": 40, "new_seo_title": 55, "tags_to_add": 40,
    "notes": 70, "status": 9, "vendor": 18, "sku": 18, "handle": 30, "product_id": 32, "url": 45,
}
WRAPPED_COLUMNS = {"tags_to_add", "notes"}

FETCH_QUERY = """
query KilnFurniture($q: String!, $after: String) {
  products(first: 50, after: $after, query: $q, sortKey: TITLE) {
    pageInfo { hasNextPage endCursor }
    nodes {
      id
      handle
      title
      status
      vendor
      productType
      tags
      publishedAt
      seo { title description }
      variants(first: 5) { nodes { sku } }
      vendorPart: metafield(namespace: "custom", key: "pdp_vendor_part_no") { value }
      size: metafield(namespace: "custom", key: "kiln_shelf_size") { value }
      shape: metafield(namespace: "custom", key: "kiln_shelf_shape") { value }
      thickness: metafield(namespace: "custom", key: "kiln_shelf_thickness") { value }
      collections(first: 20) { nodes { handle } }
    }
  }
}
"""

LIVE_QUERY = """
query LiveProducts($ids: [ID!]!) {
  nodes(ids: $ids) {
    ... on Product { id title tags seo { title } }
  }
}
"""

DEFINITIONS_QUERY = """
query ExistingDefinitions {
  metafieldDefinitions(first: 100, ownerType: PRODUCT, namespace: "custom") {
    nodes { id key validations { name value } }
  }
}
"""

CREATE_DEFINITION = """
mutation CreateDefinition($definition: MetafieldDefinitionInput!) {
  metafieldDefinitionCreate(definition: $definition) {
    createdDefinition { id namespace key }
    userErrors { field message code }
  }
}
"""

UPDATE_DEFINITION = """
mutation UpdateDefinition($definition: MetafieldDefinitionUpdateInput!) {
  metafieldDefinitionUpdate(definition: $definition) {
    updatedDefinition { id namespace key }
    userErrors { field message code }
  }
}
"""

UPDATE_PRODUCT = """
mutation UpdateProduct($product: ProductUpdateInput!) {
  productUpdate(product: $product) {
    product { id title seo { title } }
    userErrors { field message }
  }
}
"""

SET_METAFIELDS = """
mutation SetMetafields($metafields: [MetafieldsSetInput!]!) {
  metafieldsSet(metafields: $metafields) {
    metafields { id namespace key value }
    userErrors { field message code }
  }
}
"""

ADD_TAGS = """
mutation AddTags($id: ID!, $tags: [String!]!) {
  tagsAdd(id: $id, tags: $tags) {
    node { id }
    userErrors { field message }
  }
}
"""


# --------------------------------------------------------------------------
# Shopify API
# --------------------------------------------------------------------------

class Shopify:
    def __init__(self, domain, token):
        self.url = f"https://{domain}/admin/api/{API_VERSION}/graphql.json"
        self.token = token

    @classmethod
    def from_env(cls, required=True):
        domain = os.environ.get("SHOPIFY_STORE_DOMAIN")
        token = os.environ.get("SHOPIFY_ADMIN_ACCESS_TOKEN")
        if not domain or not token:
            if not required:
                return None
            raise SystemExit("Set SHOPIFY_STORE_DOMAIN and SHOPIFY_ADMIN_ACCESS_TOKEN environment variables first.")
        return cls(domain, token)

    def graphql(self, query, variables=None):
        body = json.dumps({"query": query, "variables": variables or {}}).encode()
        for attempt in range(6):
            req = urllib.request.Request(
                self.url,
                data=body,
                headers={"Content-Type": "application/json", "X-Shopify-Access-Token": self.token},
                method="POST",
            )
            try:
                with urllib.request.urlopen(req, timeout=60) as resp:
                    payload = json.loads(resp.read())
            except urllib.error.HTTPError as e:
                if e.code in (429, 502, 503) and attempt < 5:
                    time.sleep(2 ** attempt)
                    continue
                raise SystemExit(f"Shopify API error {e.code}: {e.read().decode()}")
            errors = payload.get("errors")
            if errors:
                throttled = any((err.get("extensions") or {}).get("code") == "THROTTLED" for err in errors)
                if throttled and attempt < 5:
                    time.sleep(2 ** attempt)
                    continue
                raise SystemExit(f"Shopify GraphQL errors: {errors}")
            return payload["data"]
        raise SystemExit("Shopify API still throttled after retries.")


def is_kiln_shelf(product):
    title = product["title"].lower()
    return (
        product.get("productType") == "Kiln Furniture"
        and re.search(r"shel(f|ves)", title) is not None
        and re.search(r"\bkits?\b|\bposts?\b", title) is None
    )


# --------------------------------------------------------------------------
# Title parsing
# --------------------------------------------------------------------------

# Fractions first, so "5/16" isn't read as "5" followed by "16".
DIM_RE = re.compile(r"\d+/\d+|\d+(?:\.\d+)?(?:[ -]\d+/\d+)?")


def normalize_title(title):
    t = title.replace("”", IN).replace("“", IN).replace("″", IN)
    t = re.sub(r'"\s*"', IN, t)  # 1"" and 5/8 " " -> 1" and 5/8"
    return re.sub(r"\s+", " ", t).strip()


def parse_number(token):
    m = re.fullmatch(r"(\d+(?:\.\d+)?)[ -](\d+)/(\d+)", token)
    if m:
        return Fraction(m.group(1)) + Fraction(int(m.group(2)), int(m.group(3)))
    return Fraction(token)


def fmt_inches(value):
    """Fraction/float -> '15 1/2', '5/16', '24'."""
    f = Fraction(value).limit_denominator(16)
    whole, rem = divmod(f.numerator, f.denominator)
    if rem == 0:
        return str(whole)
    frac = f"{rem}/{f.denominator}"
    return f"{whole} {frac}" if whole else frac


def parse_title(title):
    """Pull shape, plan dimensions and thickness out of a kiln shelf title."""
    low = normalize_title(title).lower()
    notes = []
    half = re.search(r"\bhalf\b", low) is not None

    sides = None
    m = re.search(r"(\d+)\s*-?\s*sided", low)
    if m:
        sides = int(m.group(1))
    elif "dodecagon" in low:
        sides = 12
    elif "decagon" in low:
        sides = 10
    elif "octagon" in low:
        sides = 8
    elif "hexagon" in low:
        sides = 6

    stripped = re.sub(r"\d+\s*-?\s*sided|\bcone\s*\d+", " ", low)
    dims = [parse_number(tok) for tok in DIM_RE.findall(stripped)]
    plan_dims = [d for d in dims if d > Fraction(3, 2)]
    thicknesses = [d for d in dims if d <= Fraction(3, 2)]

    spec = {"shape": None, "a": None, "b": None, "diameter": None, "title_size": None,
            "thickness": thicknesses[-1] if thicknesses else None, "notes": notes}
    if not spec["thickness"]:
        notes.append("No thickness in title.")

    if not plan_dims or len(plan_dims) > 2:
        notes.append(f"Could not read the shelf size from the title (found {[fmt_inches(d) for d in dims]}).")
        return spec

    if sides or re.search(r"\bround\b", low):
        kind = f"{sides}-Sided" if sides else "Round"
        spec["shape"] = f"{'Half' if half else 'Full'} {kind}"
        spec["diameter"] = max(plan_dims)
        if not half and len(plan_dims) == 2 and plan_dims[0] != plan_dims[1]:
            big, small = max(plan_dims), min(plan_dims)
            spec["title_size"] = f"{fmt_inches(big)}{IN} x {fmt_inches(small)}{IN}"
        return spec

    if len(plan_dims) == 1:
        notes.append("Only one dimension and no round/sided keyword: shape unknown.")
        return spec

    a, b = sorted(plan_dims)
    spec["a"], spec["b"] = a, b
    spec["shape"] = "Square" if a == b else "Rectangle"
    if spec["shape"] == "Rectangle" and re.search(r"\bsquare\b", low):
        notes.append(f'Old title called this {fmt_inches(b)} x {fmt_inches(a)} shelf "Square"; it is a rectangle.')
    if spec["shape"] == "Square" and re.search(r"\brect", low):
        notes.append('Old title called this shelf "Rectangle"; both sides are equal, so it is a square.')
    return spec


def product_line(product):
    vendor = product["vendor"].lower()
    low = product["title"].lower()
    if "crystolon" in low or "oxide bonded" in low:
        return "crystolon"
    if vendor == "advancer" or "advancer" in low:
        return "advancer"
    if "spectrum" in vendor or "spectrum" in low:
        return "spectrum_semi_hollow" if re.search(r"semi[- ]hollow", low) else "spectrum"
    if "cedar heights" in vendor or "corelite" in low:
        return "corelite"
    if "gillespie" in vendor or "gillespie" in low:
        return "gillespie"
    if "olympic" in vendor:
        return "olympic"
    return None


# --------------------------------------------------------------------------
# Plan
# --------------------------------------------------------------------------

def size_sort_key(value):
    """Order '12" x 24"' and '21" Diameter' values by their first number, then second."""
    nums = [parse_number(tok) for tok in DIM_RE.findall(value)]
    return (nums + [Fraction(0), Fraction(0)])[:2]


def build_row(product):
    override = OVERRIDES.get(product["handle"], {})
    spec = parse_title(product["title"])
    for field in ("shape", "a", "b", "diameter", "thickness", "title_size"):
        if field in override:
            spec[field] = override[field]
    notes = list(spec["notes"]) + list(override.get("notes", []))
    if "thickness" in override:
        notes = [n for n in notes if n != "No thickness in title."]

    line_key = override.get("line") or product_line(product)
    if not line_key:
        notes.append(f"Unknown brand/line for vendor {product['vendor']!r}; title not generated.")

    shape, thickness = spec["shape"], spec["thickness"]
    if spec["a"] is not None:
        size_value = f"{fmt_inches(spec['a'])}{IN} x {fmt_inches(spec['b'])}{IN}"
        title_size = size_value
    elif spec["diameter"] is not None:
        size_value = f"{fmt_inches(spec['diameter'])}{IN} Diameter"
        title_size = spec["title_size"] or f"{fmt_inches(spec['diameter'])}{IN}"
    else:
        size_value = title_size = ""
    thickness_value = f"{fmt_inches(thickness)}{IN}" if thickness else ""

    new_title = new_seo = ""
    if line_key and size_value and shape:
        line_title, line_seo, _ = LINES[line_key]
        new_title = f"{line_title} - {title_size} {shape}"
        dims_x = title_size
        if thickness_value:
            new_title += f", {thickness_value} Thick"
            dims_x += f" x {thickness_value}"
        new_seo = f"{line_seo} {dims_x} {shape}"
        if len(new_seo + SEO_SUFFIX) <= SEO_MAX_LEN:
            new_seo += SEO_SUFFIX
    new_title = override.get("title", new_title)
    new_seo = override.get("seo_title", new_seo)

    tags = product.get("tags") or []
    existing = {t.lower() for t in tags}
    hidden = bool(existing & HIDDEN_TAGS)
    tags_to_add = []
    if product["status"] == "ACTIVE" and product.get("publishedAt") and not hidden:
        wanted = list(CATEGORY_TAGS)
        if line_key and LINES[line_key][2]:
            wanted.append(LINES[line_key][2])
        tags_to_add = [t for t in wanted if t.lower() not in existing]
        if "kiln shelves" not in existing:
            notes.append('Missing from the "Kiln Shelves and Posts" collection until the Kiln Shelves tag is added.')
    elif hidden:
        notes.append("Hidden options-app product (OPTIONS_HIDDEN_PRODUCT): no category tags added.")
    if product["status"] == "DRAFT":
        notes.append("Draft: not visible on the storefront.")

    complete = bool(new_title and shape and size_value and thickness_value)
    approve = override.get("approve", "yes" if complete else "no")
    skus = [v["sku"] for v in product["variants"]["nodes"] if v.get("sku")]
    return {
        "approve": approve,
        "status": product["status"],
        "product_id": product["id"],
        "handle": product["handle"],
        "sku": ";".join(skus),
        "vendor": product["vendor"],
        "current_title": product["title"],
        "new_title": new_title,
        "current_seo_title": (product.get("seo") or {}).get("title") or "",
        "new_seo_title": new_seo,
        "shelf_shape": shape or "",
        "shelf_size": size_value,
        "shelf_thickness": thickness_value,
        "tags_to_add": "; ".join(tags_to_add),
        "notes": " ".join(notes),
        "url": STOREFRONT + product["handle"] if product.get("publishedAt") else "",
    }


def load_snapshot(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)["products"]


def build_plan(products):
    rows = [build_row(p) for p in products if p["status"] != "ARCHIVED" and is_kiln_shelf(p)]
    shape_rank = {s: i for i, s in enumerate(SHAPE_ORDER)}
    rows.sort(key=lambda r: (r["vendor"].lower(), shape_rank.get(r["shelf_shape"], 99),
                             size_sort_key(r["shelf_size"]), r["shelf_thickness"]))
    return rows


def facet_values(rows, field):
    values = {r[field] for r in rows if r[field] and r["approve"].strip().lower() in APPROVED}
    if field == "shelf_shape":
        return sorted(values, key=lambda v: (SHAPE_ORDER.index(v) if v in SHAPE_ORDER else 99, v))
    return sorted(values, key=size_sort_key)


def write_csv(rows, path):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=PLAN_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows, path):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("openpyxl not installed; skipped the .xlsx copy (pip install openpyxl).", file=sys.stderr)
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kiln shelves"
    ws.append(PLAN_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    review = PatternFill("solid", fgColor="FDE2E1")
    noted = PatternFill("solid", fgColor="FFF4CC")
    for row in rows:
        ws.append([row[c] for c in PLAN_COLUMNS])
        fill = review if row["approve"] != "yes" else noted if row["notes"] else None
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    for i, col in enumerate(PLAN_COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = COLUMN_WIDTHS.get(col, 15)
    for row in ws.iter_rows(min_row=2):
        for col, cell in zip(PLAN_COLUMNS, row):
            cell.alignment = Alignment(vertical="top", wrap_text=col in WRAPPED_COLUMNS)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def write_synonyms(rows, path):
    """Two-way search synonyms for Experro, so '12x24' or '24 x 12' finds titles written '12" x 24"'."""
    lines = []
    for value in facet_values(rows, "shelf_size"):
        nums = [parse_number(t) for t in DIM_RE.findall(value)]
        if "Diameter" in value or len(nums) != 2:
            continue
        # People type 10.5, not 10 1/2.
        a, b = (str(float(n)).removesuffix(".0") for n in nums)
        terms = [f"{a}x{b}", f"{a} x {b}", f"{b}x{a}", f"{b} x {a}"]
        lines.append((value.replace(IN, ""), ", ".join(dict.fromkeys(terms))))
    lines += [
        ("8-sided", "8-sided, 8 sided, octagon, octagonal"),
        ("10-sided", "10-sided, 10 sided, decagon"),
        ("12-sided", "12-sided, 12 sided, dodecagon"),
        ("silicon carbide", "silicon carbide, sic"),
        ("semi-hollow", "semi-hollow, semi hollow, hollow core"),
    ]
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["term", "synonyms"])
        writer.writerows(lines)


# --------------------------------------------------------------------------
# Apply / restore
# --------------------------------------------------------------------------

APPROVED = {"yes", "y", "true", "1"}


def read_plan(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def choices_of(definition):
    for v in definition.get("validations") or []:
        if v["name"] == "choices":
            return json.loads(v["value"])
    return []


def ensure_definitions(client, rows, confirm):
    existing = {d["key"]: d for d in client.graphql(DEFINITIONS_QUERY)["metafieldDefinitions"]["nodes"]} if client else {}
    for key, (name, description) in METAFIELDS.items():
        needed = facet_values(rows, PLAN_FIELD_FOR_KEY[key])
        if not needed:
            continue
        current = choices_of(existing[key]) if key in existing else []
        merged = needed + [c for c in current if c not in needed]
        if key == "kiln_shelf_shape":
            merged.sort(key=lambda v: (SHAPE_ORDER.index(v) if v in SHAPE_ORDER else 99, v))
        else:
            merged.sort(key=size_sort_key)
        validations = [{"name": "choices", "value": json.dumps(merged)}]
        if key not in existing:
            print(f"  create definition custom.{key} ({len(merged)} choices)")
            if confirm:
                result = client.graphql(CREATE_DEFINITION, {"definition": {
                    "name": name,
                    "namespace": "custom",
                    "key": key,
                    "description": description,
                    "type": "single_line_text_field",
                    "ownerType": "PRODUCT",
                    "validations": validations,
                    "access": {"storefront": "PUBLIC_READ"},
                    "capabilities": {"adminFilterable": {"enabled": True},
                                     "smartCollectionCondition": {"enabled": True}},
                }})["metafieldDefinitionCreate"]
                check_errors(f"create custom.{key}", result)
        elif set(merged) != set(current):
            print(f"  add choices to custom.{key}: {sorted(set(merged) - set(current), key=size_sort_key)}")
            if confirm:
                result = client.graphql(UPDATE_DEFINITION, {"definition": {
                    "namespace": "custom", "key": key, "ownerType": "PRODUCT", "validations": validations,
                }})["metafieldDefinitionUpdate"]
                check_errors(f"update custom.{key}", result)


def check_errors(label, result):
    errors = result.get("userErrors") or []
    if errors:
        raise SystemExit(f"{label} failed: {errors}")


def fetch_live(client, ids):
    live = {}
    for i in range(0, len(ids), 50):
        for node in client.graphql(LIVE_QUERY, {"ids": ids[i:i + 50]})["nodes"]:
            if node:
                live[node["id"]] = node
    return live


def cmd_fetch(args):
    client = Shopify.from_env()
    products, cursor = [], None
    while True:
        data = client.graphql(FETCH_QUERY, {"q": "product_type:'Kiln Furniture'", "after": cursor})["products"]
        products += [p for p in data["nodes"] if is_kiln_shelf(p)]
        if not data["pageInfo"]["hasNextPage"]:
            break
        cursor = data["pageInfo"]["endCursor"]
    products.sort(key=lambda p: p["title"].lower())
    snapshot = {
        "fetched_at": time.strftime("%Y-%m-%d"),
        "source": "Shopify Admin API, product_type:'Kiln Furniture', filtered to kiln shelves",
        "products": products,
    }
    os.makedirs(os.path.dirname(args.snapshot) or ".", exist_ok=True)
    with open(args.snapshot, "w", encoding="utf-8") as f:
        json.dump(snapshot, f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"Wrote {len(products)} kiln shelves to {args.snapshot}")


def cmd_plan(args):
    rows = build_plan(load_snapshot(args.snapshot))
    write_csv(rows, args.out_csv)
    write_xlsx(rows, args.out_xlsx)
    write_synonyms(rows, args.out_synonyms)
    approved = [r for r in rows if r["approve"] in APPROVED]
    print(f"Wrote {len(rows)} shelves to {args.out_csv} ({len(approved)} ready, {len(rows) - len(approved)} need review)")
    print(f"  title changes:      {sum(r['new_title'] and r['new_title'] != r['current_title'] for r in approved)}")
    print(f"  SEO title changes:  {sum(r['new_seo_title'] and r['new_seo_title'] != r['current_seo_title'] for r in approved)}")
    print(f"  shelves gaining category tags: {sum(bool(r['tags_to_add']) for r in approved)}")
    for field in ("shelf_shape", "shelf_size", "shelf_thickness"):
        print(f"  {field}: {' | '.join(facet_values(rows, field))}")
    for r in rows:
        if r["approve"] not in APPROVED:
            print(f"  NEEDS REVIEW: {r['current_title']}: {r['notes']}", file=sys.stderr)


def cmd_apply(args):
    parts = {p.strip() for p in args.only.split(",")}
    unknown = parts - {"titles", "seo", "metafields", "tags"}
    if unknown:
        raise SystemExit(f"--only: unknown part(s) {sorted(unknown)}")
    rows = [r for r in read_plan(args.plan) if r["approve"].strip().lower() in APPROVED]
    # A dry run still reads live data when credentials are set; it never writes.
    client = Shopify.from_env(required=args.confirm)
    live = fetch_live(client, [r["product_id"] for r in rows]) if client else {}
    mode = "APPLYING" if args.confirm else "DRY RUN (pass --confirm to write)"
    print(f"{mode}: {len(rows)} approved shelves, parts: {', '.join(sorted(parts))}")

    if client:
        stale = [r for r in rows if r["product_id"] in live and live[r["product_id"]]["title"] != r["current_title"]
                 and live[r["product_id"]]["title"] != r["new_title"]]
        for r in stale:
            print(f"  {'UPDATING ANYWAY' if args.force else 'SKIP'} (title changed since the plan was made): "
                  f"{live[r['product_id']]['title']}")
        if stale and not args.force:
            rows = [r for r in rows if r not in stale]
        for r in rows:
            if r["product_id"] not in live:
                print(f"  SKIP (product not found): {r['current_title']}")
        rows = [r for r in rows if r["product_id"] in live]

    if "metafields" in parts:
        ensure_definitions(client, rows, args.confirm)

    metafields = []
    for r in rows:
        pid = r["product_id"]
        now = live.get(pid) or {"title": r["current_title"], "seo": {"title": r["current_seo_title"]}}
        product = {"id": pid}
        if "titles" in parts and r["new_title"] and r["new_title"] != now["title"]:
            product["title"] = r["new_title"]
        if "seo" in parts and r["new_seo_title"] and r["new_seo_title"] != ((now.get("seo") or {}).get("title") or ""):
            product["seo"] = {"title": r["new_seo_title"]}
        if len(product) > 1:
            print(f"  update {r['current_title']!r} -> {product.get('title', '(title unchanged)')!r}")
            if args.confirm:
                result = client.graphql(UPDATE_PRODUCT, {"product": product})["productUpdate"]
                check_errors(f"productUpdate {pid}", result)
        if "metafields" in parts:
            for key, field in PLAN_FIELD_FOR_KEY.items():
                if r[field]:
                    metafields.append({"ownerId": pid, "namespace": "custom", "key": key,
                                       "type": "single_line_text_field", "value": r[field]})
        if "tags" in parts and r["tags_to_add"].strip():
            tags = [t.strip() for t in r["tags_to_add"].split(";") if t.strip()]
            if client:
                have = {t.lower() for t in live[pid]["tags"]}
                tags = [t for t in tags if t.lower() not in have]
            if tags:
                print(f"  add tags {tags} to {r['current_title']!r}")
                if args.confirm:
                    check_errors(f"tagsAdd {pid}", client.graphql(ADD_TAGS, {"id": pid, "tags": tags})["tagsAdd"])

    if metafields:
        print(f"  set {len(metafields)} size-filter metafields on {len({m['ownerId'] for m in metafields})} shelves")
        if args.confirm:
            for i in range(0, len(metafields), 25):
                result = client.graphql(SET_METAFIELDS, {"metafields": metafields[i:i + 25]})["metafieldsSet"]
                check_errors("metafieldsSet", result)
    print("Done." if args.confirm else "Dry run only; nothing was written.")


def cmd_restore(args):
    products = [p for p in load_snapshot(args.snapshot) if p["status"] != "ARCHIVED"]
    client = Shopify.from_env()
    live = fetch_live(client, [p["id"] for p in products])
    changed = 0
    for p in products:
        now = live.get(p["id"])
        if not now:
            continue
        old_seo = (p.get("seo") or {}).get("title") or ""
        now_seo = (now.get("seo") or {}).get("title") or ""
        if now["title"] == p["title"] and now_seo == old_seo:
            continue
        changed += 1
        print(f"  restore {now['title']!r} -> {p['title']!r}")
        if args.confirm:
            result = client.graphql(UPDATE_PRODUCT, {"product": {"id": p["id"], "title": p["title"],
                                                                 "seo": {"title": old_seo}}})["productUpdate"]
            check_errors(f"restore {p['id']}", result)
    print(f"{changed} product(s) {'restored' if args.confirm else 'would be restored (pass --confirm)'}."
          " Metafields and added tags are left in place.")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("fetch", help="Download kiln shelves from Shopify into the snapshot file.")
    p.add_argument("--snapshot", default=DEFAULT_SNAPSHOT)
    p.set_defaults(func=cmd_fetch)

    p = sub.add_parser("plan", help="Build the review sheet from the snapshot (no Shopify writes).")
    p.add_argument("--snapshot", default=DEFAULT_SNAPSHOT)
    p.add_argument("--out-csv", default=DEFAULT_PLAN_CSV)
    p.add_argument("--out-xlsx", default=DEFAULT_PLAN_XLSX)
    p.add_argument("--out-synonyms", default=DEFAULT_SYNONYMS)
    p.set_defaults(func=cmd_plan)

    p = sub.add_parser("apply", help="Write approved plan rows to Shopify (dry run without --confirm).")
    p.add_argument("--plan", default=DEFAULT_PLAN_CSV)
    p.add_argument("--only", default="titles,seo,metafields,tags",
                   help="Comma-separated parts to apply: titles, seo, metafields, tags (default: all).")
    p.add_argument("--confirm", action="store_true", help="Actually write to Shopify.")
    p.add_argument("--force", action="store_true", help="Also update products whose title changed since the plan.")
    p.set_defaults(func=cmd_apply)

    p = sub.add_parser("restore", help="Restore titles and SEO titles from the snapshot (dry run without --confirm).")
    p.add_argument("--snapshot", default=DEFAULT_SNAPSHOT)
    p.add_argument("--confirm", action="store_true")
    p.set_defaults(func=cmd_restore)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
